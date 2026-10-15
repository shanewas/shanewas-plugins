"""Excel COM recalc. Windows + MS Excel only. Opens file in real Excel, full rebuild, save, scan for errors."""
import time
from pathlib import Path

import openpyxl

ERROR_STRINGS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!")


def recalc(path: str, timeout: int = 60) -> dict:
    p = Path(path).resolve()
    win_path = str(p)  # COM needs native Windows backslash path; Path.resolve() on Windows already gives backslashes

    import win32com.client

    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    wb = None
    try:
        wb = excel.Workbooks.Open(win_path)
        excel.CalculateFullRebuild()
        deadline = time.time() + timeout
        while excel.CalculationState != 0:  # 0 == xlDone
            if time.time() > deadline:
                raise TimeoutError(f"recalc did not finish within {timeout}s")
            time.sleep(0.2)
        wb.Save()
        wb.Close(SaveChanges=False)
        wb = None
    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        excel.Quit()
        del excel

    return _scan_errors(p)


def _scan_errors(p: Path) -> dict:
    wb = openpyxl.load_workbook(p, data_only=True)
    total_formulas = 0
    total_errors = 0
    error_summary: dict[str, dict] = {}

    wb_formulas = openpyxl.load_workbook(p, data_only=False)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws_f = wb_formulas[sheet_name]
        for row in ws_f.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    total_formulas += 1
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v in ERROR_STRINGS:
                    total_errors += 1
                    entry = error_summary.setdefault(v, {"count": 0, "locations": []})
                    entry["count"] += 1
                    if len(entry["locations"]) < 20:
                        entry["locations"].append(f"{sheet_name}!{cell.coordinate}")
    wb.close()
    wb_formulas.close()

    result = {
        "status": "errors_found" if total_errors else "success",
        "total_errors": total_errors,
        "total_formulas": total_formulas,
    }
    if total_errors:
        result["error_summary"] = error_summary
    return result
